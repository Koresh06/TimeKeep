from io import BytesIO
from docx import Document
from docx.shared import Pt
from pathlib import Path

TEMPLATE_PATH = Path("src/api/v1/day_off/reports/template/report_template.docx")

class ReportGenerator:
    def __init__(self, data: dict, template_path: Path = TEMPLATE_PATH):
        """
        Инициализация генератора отчётов.

        :param template_path: Путь к шаблону .docx
        :param data: Данные для замены в шаблоне
        """
        self.template_path = template_path
        self.data = data
        self.doc = Document(template_path)


    def format_overtimes(self, info_overtimes):
        """Форматирование списка переработок в строку"""
        return "; ".join([f"{item['description']} - {item['hours']} ч." for item in info_overtimes])


    def set_font(self, run, font_name="Times New Roman", font_size=15):
        """Устанавливает шрифт и размер текста"""
        run.font.name = font_name
        run.font.size = Pt(font_size)


    def replace_placeholders(self):
        """Заменяет шаблонные переменные в документе"""
        for paragraph in self.doc.paragraphs:
            for key, value in self.data.items():
                if key in paragraph.text:
                    if isinstance(value, list):
                        value = self.format_overtimes(value)  # Преобразуем список в строку
                    paragraph.text = paragraph.text.replace(key, value)

                    # Применяем шрифт ко всем частям текста (runs)
                    for run in paragraph.runs:
                        self.set_font(run)


    def get_report_bytes(self):
        """Создаёт документ и возвращает его как байтовый поток"""
        self.replace_placeholders()
        file_stream = BytesIO()
        self.doc.save(file_stream)
        file_stream.seek(0)  # Перемещаем указатель в начало файла
        return file_stream
    
